"""Regression tests for the ppt-generator renderer and validator.

Run from the skill root:
    python3 -m unittest discover -s tests -v
"""
import copy
import io
import json
import os
import sys
import unittest

sys.dont_write_bytecode = True
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import build_pptx  # noqa: E402
import textfit  # noqa: E402
import validate_spec  # noqa: E402
import icons  # noqa: E402
import build_icon_set  # noqa: E402
import render_preview  # noqa: E402
import audit_pptx  # noqa: E402
import inspect_pptx  # noqa: E402
from templates import THEMES  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

EXAMPLE = os.path.join(ROOT, "assets", "spec.example.json")
KOREAN = os.path.join(ROOT, "tests", "fixtures", "korean.spec.json")
ASSERTION = os.path.join(ROOT, "tests", "fixtures", "assertion.spec.json")
LAYOUTS_FX = os.path.join(ROOT, "tests", "fixtures", "layouts.spec.json")


def long_bullets(n=14):
    return {"slides": [{"layout": "bullets", "title": "점검 항목", "numbered": True,
                        "subtitle": "부제목", "notes": "노트",
                        "bullets": ["항목 %d: 접근 권한 정기 검토 및 감사 로그 확인" % i
                                    for i in range(1, n + 1)]}]}


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def render(spec):
    buf = io.BytesIO()
    build_pptx.build(spec, buf, report=False)
    buf.seek(0)
    return buf


class BuildTests(unittest.TestCase):
    def test_every_template_builds_example_without_overflow(self):
        spec = load(EXAMPLE)
        for name in THEMES:
            with self.subTest(template=name):
                spec["template"] = name
                prs = Presentation(render(spec))
                self.assertEqual(len(prs.slides), len(spec["slides"]))
                report = build_pptx.fit_report()
                self.assertFalse([l for l in report if l.startswith(("OVERFLOW", "WARN"))],
                                 report)

    def test_korean_fixture_has_no_overflow_or_bounds_issue(self):
        for name in THEMES:
            with self.subTest(template=name):
                spec = load(KOREAN)
                spec["template"] = name
                render(spec)
                report = build_pptx.fit_report()
                self.assertFalse([l for l in report if l.startswith(("OVERFLOW", "WARN"))],
                                 report)

    def test_korean_runs_get_east_asian_font_and_language(self):
        prs = Presentation(render(load(KOREAN)))
        hangul_runs = 0
        for slide in prs.slides:
            for shp in slide.shapes:
                if not shp.has_text_frame:
                    continue
                for p in shp.text_frame.paragraphs:
                    for r in p.runs:
                        if not textfit.has_hangul(r.text):
                            continue
                        hangul_runs += 1
                        rPr = r._r.find(qn("a:rPr"))
                        ea = rPr.find(qn("a:ea"))
                        self.assertIsNotNone(ea, r.text)
                        self.assertEqual(ea.get("typeface"), "Malgun Gothic")
                        self.assertEqual(rPr.get("lang"), "ko-KR")
                        # a:ea must follow a:latin (schema order)
                        tags = [c.tag for c in rPr]
                        self.assertLess(tags.index(qn("a:latin")), tags.index(qn("a:ea")))
                        self.assertEqual(p._p.pPr.get("eaLnBrk"), "0")
                        self.assertEqual(p._p.pPr.get("latinLnBrk"), "0")
        self.assertGreater(hangul_runs, 20)

    def test_theme_fonts_written(self):
        buf = render(load(KOREAN))
        with zipfile.ZipFile(buf) as z:
            theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
        self.assertIn('<a:ea typeface="Malgun Gothic"/>', theme)

    def test_bullets_are_real_paragraph_bullets(self):
        prs = Presentation(render(load(EXAMPLE)))
        bullets = 0
        for slide in prs.slides:
            for shp in slide.shapes:
                if not shp.has_text_frame:
                    continue
                for p in shp.text_frame.paragraphs:
                    pPr = p._p.pPr
                    if pPr is not None and pPr.find(qn("a:buChar")) is not None:
                        bullets += 1
                        self.assertFalse(p.text.startswith(("●", "–  ", "·  ")), p.text)
                        self.assertLess(int(pPr.get("indent")), 0)  # hanging indent
        self.assertGreater(bullets, 10)

    def test_chart_series_use_theme_palette(self):
        spec = load(EXAMPLE)
        for name in ("corporate", "dark"):
            with self.subTest(template=name):
                spec["template"] = name
                theme = THEMES[name]
                prs = Presentation(render(spec))
                charts = [s.chart for sl in prs.slides for s in sl.shapes if s.has_chart]
                self.assertTrue(charts)
                first = charts[0].plots[0].series[0]
                palette = theme.get("chart_colors") or [theme["primary"]]
                self.assertEqual(str(first.format.fill.fore_color.rgb), palette[0])
                # axis/legend text must contrast with the slide background
                txt = str(charts[0].font.color.rgb)
                self.assertGreaterEqual(
                    validate_spec.contrast(txt, theme["background"]), 4.5)

    def test_highlight_greys_out_other_bars(self):
        spec = {"slides": [{"layout": "chart", "title": "Mix", "chart": {
            "type": "column", "categories": ["A", "B", "C"], "series": [3, 2, 1],
            "highlight": "B"}}]}
        prs = Presentation(render(spec))
        ser = [s for s in prs.slides[0].shapes if s.has_chart][0].chart.plots[0].series[0]
        colors = [str(ser.points[i].format.fill.fore_color.rgb) for i in range(3)]
        self.assertEqual(colors[1], THEMES[build_pptx.DEFAULT_THEME]["accent"])
        self.assertEqual(colors[0], colors[2])
        self.assertNotEqual(colors[0], colors[1])

    def test_speaker_notes(self):
        prs = Presentation(render(load(EXAMPLE)))
        noted = [s for s in prs.slides
                 if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()]
        self.assertEqual(len(noted), sum(1 for s in load(EXAMPLE)["slides"] if s.get("notes")))

    def test_long_title_moves_body_down(self):
        base = {"slides": [{"layout": "bullets", "title": "Short", "bullets": ["a"]}]}
        long = copy.deepcopy(base)
        long["slides"][0]["title"] = ("A very long slide title that certainly cannot fit "
                                      "on a single line at the heading size of the theme")

        def body_top(spec):
            slide = Presentation(render(spec)).slides[0]
            return max(s.top for s in slide.shapes
                       if s.has_text_frame and s.text_frame.text == "a")
        self.assertGreater(body_top(long), body_top(base))


    def test_korean_text_is_never_italic(self):
        spec = {"slides": [{"layout": "bullets", "title": "제목",
                            "subtitle": "4개 분기 연속 상승, 4Q 61% 도달",
                            "bullets": ["항목"]}]}
        slide = Presentation(render(spec)).slides[0]
        runs = [r for s in slide.shapes if s.has_text_frame
                for p in s.text_frame.paragraphs for r in p.runs
                if "분기" in r.text]
        self.assertTrue(runs)
        self.assertFalse(any(r.font.italic for r in runs))

    def test_orphaned_title_word_is_reported(self):
        spec = {"slides": [{"layout": "title",
                            "title": "생성형 AI 고객지원 챗봇 도입 1년 성과"}]}
        render(spec)
        self.assertTrue(any(l.startswith("ORPHAN") for l in build_pptx.fit_report()))


    def _title_fit(self, style, overflow):
        """Fit a title whose full-size width is `overflow` x the line width."""
        text = " ".join(["word"] * 40)
        size = 28
        width_pt = (textfit.text_width_pt(text, size, None, bold=True)
                    / (overflow * textfit.WIDTH_SLACK))
        build_pptx._CTX["title_style"] = style
        try:
            return build_pptx._title_lines_and_size(text, width_pt * 12700, size, None)
        finally:
            build_pptx._CTX["title_style"] = "topic"

    def test_assertion_titles_wrap_at_full_size_topic_titles_shrink(self):
        # 1.22 lines at full size: topic shrinks to one line, assertion wraps.
        size, lines = self._title_fit("topic", 1.22)
        self.assertEqual(lines, 1)
        self.assertLess(size, 28)
        self.assertEqual(self._title_fit("assertion", 1.22), (28, 2))

    def test_assertion_title_shrinks_slightly_to_avoid_orphan(self):
        # 1.06 lines would strand one word on line two: shrink 10% instead.
        size, lines = self._title_fit("assertion", 1.06)
        self.assertEqual(lines, 1)
        self.assertAlmostEqual(size, 28 * 0.9)

    def test_assertion_fixture_renders_clean(self):
        render(load(ASSERTION))
        report = build_pptx.fit_report()
        self.assertFalse([l for l in report if l.startswith(("OVERFLOW", "WARN"))],
                         report)


    def test_numbered_bullets_use_autonum_with_start(self):
        spec = {"slides": [{"layout": "bullets", "title": "Steps", "numbered": True,
                            "start_at": 3, "bullets": ["a", {"text": "sub", "level": 1}, "b"]}]}
        slide = Presentation(render(spec)).slides[0]
        nums = []
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            for p in shp.text_frame.paragraphs:
                an = p._p.pPr.find(qn("a:buAutoNum")) if p._p.pPr is not None else None
                if an is not None:
                    nums.append((p.text, an.get("type"), an.get("startAt")))
        self.assertEqual(nums, [("a", "arabicPeriod", "3"), ("b", "arabicPeriod", "4")])

    def test_long_bullets_split_with_continuation(self):
        prs = Presentation(render(long_bullets()))
        self.assertGreater(len(prs.slides), 1)
        self.assertTrue(any(l.startswith("SPLIT") for l in build_pptx.fit_report()))
        self.assertFalse([l for l in build_pptx.fit_report()
                          if l.startswith(("OVERFLOW", "FIT"))])
        texts = [[sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
                 for sl in prs.slides]
        self.assertTrue(any("점검 항목 (계속)" in t for t in texts[1]))
        # numbering continues on the continuation slide
        starts = []
        for sl in prs.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        if p._p.pPr is not None and p._p.pPr.find(qn("a:buAutoNum")) is not None:
                            starts.append(int(p._p.pPr.find(qn("a:buAutoNum")).get("startAt")))
        self.assertEqual(starts, list(range(1, 15)))
        # notes and subtitle stay on the first part only
        self.assertTrue(prs.slides[0].has_notes_slide)
        self.assertFalse(any("부제목" in t for t in texts[1]))

    def test_split_parts_are_balanced(self):
        prs = Presentation(render(long_bullets(14)))
        counts = []
        for sl in prs.slides:
            counts.append(sum(1 for sh in sl.shapes if sh.has_text_frame
                              for p in sh.text_frame.paragraphs
                              if p._p.pPr is not None
                              and p._p.pPr.find(qn("a:buAutoNum")) is not None))
        self.assertLessEqual(max(counts) - min(counts), 1)

    def test_split_can_be_disabled(self):
        spec = long_bullets()
        spec["slides"][0]["split"] = False
        prs = Presentation(render(spec))
        self.assertEqual(len(prs.slides), 1)

    def test_long_table_repeats_header(self):
        rows = [["A%d" % i, "$%d" % i, "Owner %d" % i] for i in range(20)]
        spec = {"slides": [{"layout": "table", "title": "Accounts",
                            "table": {"header": ["Account", "ARR", "Owner"], "rows": rows}}]}
        prs = Presentation(render(spec))
        tables = [sh.table for sl in prs.slides for sh in sl.shapes if sh.has_table]
        self.assertGreater(len(tables), 1)
        for t in tables:
            self.assertEqual(t.cell(0, 0).text, "Account")
        self.assertEqual(sum(len(t.rows) - 1 for t in tables), 20)
        titles = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
        self.assertIn("Accounts (cont.)", titles)

    def test_sparse_bullets_slide_is_reported(self):
        render({"slides": [{"layout": "bullets", "title": "Plan",
                            "bullets": ["One", "Two", "Three"]}]})
        self.assertTrue(any(l.startswith("SPARSE") for l in build_pptx.fit_report()))
        render({"slides": [{"layout": "bullets", "title": "Plan",
                            "bullets": ["Point number %d with some detail" % i
                                        for i in range(6)]}]})
        self.assertFalse(any(l.startswith("SPARSE") for l in build_pptx.fit_report()))

    def test_new_layouts_build_clean_in_every_template(self):
        for name in THEMES:
            with self.subTest(template=name):
                spec = load(LAYOUTS_FX)
                spec["template"] = name
                prs = Presentation(render(spec))
                self.assertEqual(len(prs.slides), len(spec["slides"]))
                report = build_pptx.fit_report()
                self.assertFalse([l for l in report
                                  if l.startswith(("OVERFLOW", "WARN", "FIT"))], report)


class IconTests(unittest.TestCase):
    def test_names_and_aliases_resolve(self):
        self.assertEqual(icons.resolve("lightbulb"), "lightbulb")
        self.assertEqual(icons.resolve("Shield Check"), "shield-check")
        self.assertEqual(icons.resolve("보안"), "shield-check")
        self.assertEqual(icons.resolve("idea"), "lightbulb")
        self.assertIsNone(icons.resolve("no-such-icon"))
        self.assertGreaterEqual(len(icons.names()), 150)
        for alias, target in icons.ALIASES.items():
            self.assertIn(target, icons.names(), alias)

    def test_icon_is_native_custgeom_shape(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        sp = icons.add_icon(slide, "rocket", 0, 0, 914400, "FF0000")
        geom = sp._element.spPr.find(qn("a:custGeom"))
        self.assertIsNotNone(geom)
        self.assertIsNone(sp._element.spPr.find(qn("a:prstGeom")))
        self.assertGreater(len(geom.findall(".//" + qn("a:cubicBezTo"))), 3)
        self.assertEqual(str(sp.fill.fore_color.rgb), "FF0000")
        self.assertIsNone(icons.add_icon(slide, "no-such-icon", 0, 0, 914400, "FF0000"))

    def test_arc_converts_to_cubics_ending_at_endpoint(self):
        ops = build_icon_set.parse_path("M0,0 A10,10 0 0 1 20,0 Z")
        cubics = [o for o in ops if o[0] == "C"]
        self.assertEqual(len(cubics), 2)                 # 180 degrees -> 2 x 90
        self.assertAlmostEqual(cubics[-1][5], 20, places=6)
        self.assertAlmostEqual(cubics[-1][6], 0, places=6)
        # sweep=1 from (0,0) to (20,0) bulges toward -y in SVG coordinates
        self.assertLess(min(c[2] for c in cubics), -5)

    def test_relative_commands_and_shorthands(self):
        ops = build_icon_set.parse_path("m10,10 h5 v5 l-5,0 z")
        self.assertEqual([o[0] for o in ops], ["M", "L", "L", "L", "Z"])
        self.assertEqual(ops[2][1:], (15.0, 15.0))

    def test_icon_grid_unknown_icon_falls_back_to_number(self):
        spec = {"slides": [{"layout": "icon_grid", "title": "x", "items": [
            {"icon": "target", "title": "A"}, {"icon": "zzz", "title": "B"},
            {"icon": "lightning", "title": "C"}]}]}
        slide = Presentation(render(spec)).slides[0]
        names = [sh.name for sh in slide.shapes]
        self.assertIn("Icon: target", names)
        self.assertIn("Icon: lightning", names)
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        self.assertIn("2", texts)
        _, warnings = validate_spec.validate(spec)
        self.assertTrue(any("unknown icon 'zzz'" in w for w in warnings))


class LayoutMetricsTests(unittest.TestCase):
    def _png(self, path, ink_until):
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (400, 225), "white")
        d = ImageDraw.Draw(im)
        d.rectangle((0, 0, 400, 25), fill="black")           # header bar
        d.rectangle((30, 40, 300, int(225 * ink_until)), fill="navy")
        im.save(path)

    def test_empty_bottom_detected_despite_dark_header(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            sparse, full = os.path.join(d, "slide-01.png"), os.path.join(d, "slide-02.png")
            self._png(sparse, 0.45)
            self._png(full, 0.85)
            m = render_preview.layout_metrics([sparse, full])
        self.assertGreater(m[0]["empty_bottom"], render_preview.EMPTY_BOTTOM_LIMIT)
        self.assertLess(m[1]["empty_bottom"], render_preview.EMPTY_BOTTOM_LIMIT)


def _foreign_deck(path):
    """A deck as another tool might hand it over: default placeholders, a
    group, tiny and low-contrast text, a stretched low-res picture, an
    off-slide box, Hangul without a:ea."""
    from PIL import Image
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    img = path + ".png"
    Image.new("RGB", (200, 100), "#3366CC").save(img)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "분기 실적 보고"
    s.placeholders[1].text_frame.text = "매출 증가"
    tb = s.shapes.add_textbox(Inches(1), Inches(5.6), Inches(7), Inches(0.4))
    tb.text_frame.text = "출처: 내부 데이터"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    tb2 = s.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(8), Inches(0.4))
    tb2.text_frame.text = "연한 회색 주석"
    tb2.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    tb2.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_picture(img, Inches(0.5), Inches(0.5), width=Inches(8), height=Inches(6))
    grp = s2.shapes.add_group_shape()
    grp.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5)).text_frame.text = "그룹 A"
    grp.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(3), Inches(0.5)).text_frame.text = "그룹 B"
    s2.shapes.add_textbox(Inches(9), Inches(3), Inches(3), Inches(1)).text_frame.text = "off slide"
    prs.save(path)


class AuditTests(unittest.TestCase):
    def test_foreign_deck_flaws_are_found(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "foreign.pptx")
            _foreign_deck(path)
            f = audit_pptx.audit(path)
        cats = {(x["slide"], x["category"]) for x in f}
        self.assertIn((1, "Inappropriate Font Sizing"), cats)
        self.assertIn((1, "Insufficient Contrast"), cats)
        self.assertIn((2, "Content Overflow or Truncation"), cats)      # off-slide
        self.assertIn((2, "Low-Quality or Distorted Images"), cats)
        self.assertIn((2, "Improper Image Scaling or Size"), cats)
        self.assertIn((2, "Occluded Content"), cats)                    # inside a group
        self.assertIn((2, "Accessibility"), cats)                       # file-name alt text
        self.assertTrue(any(x["slide"] == 0 and "Korean" in x["message"] for x in f))

    def test_generated_decks_audit_clean_in_every_template(self):
        import tempfile
        for name in THEMES:
            with self.subTest(template=name), tempfile.TemporaryDirectory() as d:
                spec = load(EXAMPLE)
                spec["template"] = name
                path = os.path.join(d, "deck.pptx")
                build_pptx.build(spec, path, report=False)
                bad = [x for x in audit_pptx.audit(path) if x["severity"] != "info"]
                self.assertEqual(bad, [])

    def test_overflowing_deck_is_an_error(self):
        import tempfile
        spec = {"slides": [{"layout": "bullets", "title": "Too much", "split": False,
                            "bullets": ["항목 %d번은 꽤 길게 작성된 설명 문장입니다" % i
                                        for i in range(30)]}]}
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "deck.pptx")
            build_pptx.build(spec, path, report=False)
            f = audit_pptx.audit(path)
        self.assertTrue(any(x["severity"] == "error"
                            and x["category"] == "Content Overflow or Truncation" for x in f))

    def test_image_slides_carry_alt_text(self):
        import tempfile
        from PIL import Image
        with tempfile.TemporaryDirectory() as d:
            img = os.path.join(d, "a.png")
            Image.new("RGB", (1600, 900), "white").save(img)
            spec = {"slides": [{"layout": "image", "title": "Plant layout", "image": img},
                               {"layout": "image", "title": "x", "image": img,
                                "alt": "Floor plan of plant 2", "position": "full"}]}
            prs = Presentation(render(spec))
        descr = [sh._element.nvPicPr.cNvPr.get("descr") for sl in prs.slides
                 for sh in sl.shapes if sh.shape_type == 13]
        self.assertEqual(descr, ["Plant layout", "Floor plan of plant 2"])

    def test_section_number_stays_readable_on_dark_bleed(self):
        spec = {"template": "standard", "slides": [
            {"layout": "section", "number": "01", "title": "Part"}]}
        slide = Presentation(render(spec)).slides[0]
        run = [r for sh in slide.shapes if sh.has_text_frame
               for p in sh.text_frame.paragraphs for r in p.runs if r.text == "01"][0]
        bleed = THEMES["standard"].get("bleed", THEMES["standard"]["primary"])
        self.assertGreaterEqual(validate_spec.contrast(str(run.font.color.rgb), bleed), 3.0)

    def test_inspect_geometry(self):
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "deck.pptx")
            build_pptx.build(load(EXAMPLE), path, report=False)
            slides = inspect_pptx.inspect(path, geometry=True)
        shapes = [sh for s in slides for sh in s["shapes"] if "text" in sh]
        self.assertTrue(all(len(sh["box_in"]) == 4 for sh in shapes))
        self.assertTrue(any("font_pt" in sh for sh in shapes))


class ValidatorTests(unittest.TestCase):
    def test_example_is_clean(self):
        errors, warnings = validate_spec.validate(load(EXAMPLE))
        e2, w2 = validate_spec.dry_run(load(EXAMPLE))
        self.assertEqual(errors + e2, [])
        self.assertEqual(warnings + w2, [])

    def test_overflow_is_an_error_when_split_is_off(self):
        spec = {"slides": [{"layout": "bullets", "title": "Too much", "split": False,
                            "bullets": ["항목 %d번은 꽤 길게 작성된 설명 문장입니다" % i
                                        for i in range(22)]}]}
        errors, _ = validate_spec.dry_run(spec)
        self.assertTrue(any("OVERFLOW" in e for e in errors), errors)

    def test_overlong_list_is_split_not_failed(self):
        spec = {"slides": [{"layout": "bullets", "title": "Too much",
                            "bullets": ["항목 %d번은 꽤 길게 작성된 설명 문장입니다" % i
                                        for i in range(22)]}]}
        errors, warnings = validate_spec.dry_run(spec)
        self.assertEqual(errors, [])
        self.assertTrue(any(w.startswith("SPLIT") for w in warnings), warnings)

    def test_unsplittable_single_bullet_still_overflows(self):
        spec = {"slides": [{"layout": "bullets", "title": "One giant bullet",
                            "bullets": ["아주 긴 문장 " * 300]}]}
        errors, _ = validate_spec.dry_run(spec)
        self.assertTrue(any("OVERFLOW" in e for e in errors), errors)

    def test_chart_length_mismatch_and_bad_highlight(self):
        spec = {"slides": [{"layout": "chart", "title": "x", "chart": {
            "categories": ["A", "B"], "series": [1, 2, 3], "highlight": "Z"}}]}
        errors, _ = validate_spec.validate(spec)
        self.assertTrue(any("3 values for 2 categories" in e for e in errors))
        self.assertTrue(any("highlight" in e for e in errors))

    def test_non_numeric_chart_values_rejected(self):
        spec = {"slides": [{"layout": "chart", "title": "x", "chart": {
            "categories": ["A"], "series": ["$3M"]}}]}
        errors, _ = validate_spec.validate(spec)
        self.assertTrue(any("non-numeric" in e for e in errors))

    def test_density_and_rhythm_warnings(self):
        spec = {"slides": [{"layout": "bullets", "title": "t%d" % i,
                            "bullets": ["x"] * (8 if i == 0 else 1)} for i in range(3)]}
        _, warnings = validate_spec.validate(spec)
        self.assertTrue(any("top-level bullets" in w for w in warnings))
        self.assertTrue(any("in a row" in w for w in warnings))

    def test_assertion_fixture_has_no_title_warnings(self):
        errors, warnings = validate_spec.validate(load(ASSERTION))
        self.assertEqual(errors, [])
        self.assertEqual([w for w in warnings if "title" in w], [])

    def test_label_titles_flagged_in_assertion_deck(self):
        spec = {"title_style": "assertion", "slides": [
            {"layout": "bullets", "title": "분기별 자동 해결률", "bullets": ["a"]},
            {"layout": "bullets", "title": "보고 순서", "title_style": "topic",
             "bullets": ["a"]},
            {"layout": "section", "title": "성과 상세"}]}
        _, warnings = validate_spec.validate(spec)
        flagged = [w for w in warnings if "is a label" in w]
        self.assertEqual(len(flagged), 1)
        self.assertIn("slide 1", flagged[0])

    def test_sentence_titles_flagged_in_topic_deck(self):
        spec = {"slides": [
            {"layout": "bullets", "title": "매출이 18% 증가했다", "bullets": ["a"]},
            {"layout": "bullets", "title": "Revenue grew 18%.", "bullets": ["a"]},
            {"layout": "bullets", "title": "18% Revenue Growth", "bullets": ["a"]}]}
        _, warnings = validate_spec.validate(spec)
        self.assertEqual(len([w for w in warnings if "reads like a sentence" in w]), 2)

    def test_unknown_title_style_is_an_error(self):
        errors, _ = validate_spec.validate(
            {"title_style": "headline", "slides": [{"layout": "title", "title": "x"}]})
        self.assertTrue(any("title_style" in e for e in errors))

    def test_new_layout_rules(self):
        spec = {"slides": [
            {"layout": "matrix", "title": "m", "quadrants": [{}, {}, {}], "highlight": 5},
            {"layout": "agenda", "title": "a", "items": ["x", "y"], "current": 3},
            {"layout": "big_number", "title": "b"},
            {"layout": "cards", "title": "c", "cards": [{"title": "only one"}]}]}
        errors, warnings = validate_spec.validate(spec)
        self.assertTrue(any("exactly 4 quadrants" in e for e in errors))
        self.assertTrue(any("highlight must be 1–4" in e for e in errors))
        self.assertTrue(any("current must be" in e for e in errors))
        self.assertTrue(any("'value'" in e for e in errors))
        self.assertTrue(any("cards" in w for w in warnings))

    def test_layouts_fixture_is_clean(self):
        errors, warnings = validate_spec.validate(load(LAYOUTS_FX))
        self.assertEqual((errors, warnings), ([], []))

    def test_all_templates_pass_contrast(self):
        for name, theme in THEMES.items():
            with self.subTest(template=name):
                warnings = []
                validate_spec.lint_theme(theme, warnings)
                self.assertEqual(warnings, [])


class TextFitTests(unittest.TestCase):
    def test_hangul_is_wider_than_latin(self):
        self.assertGreater(textfit.text_width_pt("가나다라", 18),
                           textfit.text_width_pt("abcd", 18) * 1.6)

    def test_wrap_counts_lines(self):
        self.assertEqual(textfit.wrap_line_count("short", 500, 18), 1)
        self.assertGreater(textfit.wrap_line_count("word " * 60, 300, 18), 3)

    def test_fit_scale_shrinks_then_gives_up(self):
        paras = [{"text": "word " * 40, "size": 18}]
        scale, fits = textfit.fit_scale(paras, 400, 90)
        self.assertTrue(fits)
        self.assertLess(scale, 1.0)
        scale, fits = textfit.fit_scale(paras * 20, 400, 120)
        self.assertFalse(fits)


if __name__ == "__main__":
    unittest.main()
